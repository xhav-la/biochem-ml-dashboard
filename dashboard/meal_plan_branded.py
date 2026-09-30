# -*- coding: utf-8 -*-
"""
meal_plan_branded.py
======================
Dizajni i BRENDUAR i planit ushqimor (logo "Nora Limani Bektashi · Pri
Nutrition"), për Excel dhe PDF. Paleta ndjek ngjyrën e logos (ulliri
#606D3E); fontet e PDF-së (Gilda Display + Lato, licencë OFL) ndodhen
te dashboard/assets/fonts, që dokumenti të dalë njësoj kudo.

Dizajni klasik (si modeli origjinal Excel) mbetet te meal_plan_export.py.
"""
import io
import os
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from meal_plan_export import DAYS, MEALS, _est_lines, fmt_date, fmt_num

ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")
LOGO = os.path.join(ASSETS, "logo_nlb.png")  # vetëm monogrami NLB + "PRI NUTRITION"
FONTS = os.path.join(ASSETS, "fonts")

# ---- Paleta e markës (nga logo) ----
OLIVE = "606D3E"
OLIVE_DARK = "4A5530"
SAGE = "DDE2CF"
CREAM = "F7F5EE"
LINE = "C9CFB6"
TEXT = "2F3322"
MUTED = "7A7F6A"

BRAND_LINE = "NORA LIMANI BEKTASHI  ·  PRI NUTRITION"


def _bmi(plan):
    p, g = plan.get("pesha"), plan.get("gjatesia")
    if p and g:
        return f"{p / ((g / 100) ** 2):.1f}"
    return "..."


def _period(plan):
    return f"{fmt_date(plan['data_fillimit'])} – {fmt_date(plan['data_mbarimit'])}"


# =================================================================
# EXCEL
# =================================================================
def build_branded_xlsx(plan) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Plani ushqimor"
    ws.sheet_view.showGridLines = False

    H_FONT = "Georgia"   # serif elegant, i instaluar në Windows/Mac
    B_FONT = "Calibri"   # tekst i pastër, i lexueshëm
    thin = Side(style="thin", color=LINE)
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)

    widths = {"A": 16, "B": 3, "C": 40, "D": 22, "E": 44, "F": 22, "G": 40, "H": 24, "I": 32}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    meal_cols = ["C", "D", "E", "F", "G", "H"]

    # ---- Koka: logo majtas, titulli djathtas ----
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 8
    if os.path.exists(LOGO):
        img = XLImage(LOGO)
        ratio = img.width / img.height          # proporcionet origjinale të imazhit
        img.height = 84
        img.width = int(84 * ratio)
        ws.add_image(img, "A1")

    ws.merge_cells("F1:I1")
    ws["F1"] = "PLAN USHQIMOR"
    ws["F1"].font = Font(name=H_FONT, size=24, color=OLIVE)
    ws["F1"].alignment = Alignment(horizontal="right", vertical="bottom")
    ws.merge_cells("F2:I2")
    ws["F2"] = f"{plan['titulli']}   ·   {_period(plan)}"
    ws["F2"].font = Font(name=B_FONT, size=12, bold=True, color=MUTED)
    ws["F2"].alignment = Alignment(horizontal="right", vertical="top")
    for col in "ABCDEFGHI":
        ws[f"{col}3"].border = Border(bottom=Side(style="medium", color=OLIVE))

    # ---- Të dhënat e klientit (rreshti 5) ----
    ws.row_dimensions[4].height = 10
    ws.row_dimensions[5].height = 16
    ws.row_dimensions[6].height = 22
    info = [("A", "KLIENTI", plan["emri"], "D"), ("E", "PESHA", fmt_num(plan["pesha"], "kg"), None),
            ("F", "GJATËSIA", fmt_num(plan["gjatesia"], "cm"), None), ("G", "BMI", _bmi(plan), None)]
    for col, label, value, merge_to in info:
        if merge_to:
            ws.merge_cells(f"{col}5:{merge_to}5")
            ws.merge_cells(f"{col}6:{merge_to}6")
        ws[f"{col}5"] = label
        ws[f"{col}5"].font = Font(name=B_FONT, size=9, bold=True, color=MUTED)
        ws[f"{col}6"] = value
        ws[f"{col}6"].font = Font(name=H_FONT, size=14, color=TEXT)
        for c in (ws[f"{col}5"], ws[f"{col}6"]):
            c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[7].height = 12

    # ---- Koka e tabelës (rreshti 8) ----
    HR = 8
    ws.row_dimensions[HR].height = 38
    head_fill = PatternFill("solid", fgColor=OLIVE)
    ws.merge_cells(f"A{HR}:B{HR}")
    ws[f"A{HR}"] = "DITA"
    headers = {"A": "DITA", "I": "LISTA USHQIMORE"}
    for col, (key, name, _) in zip(meal_cols, MEALS):
        headers[col] = f"{name.upper()}\n{plan['oraret'].get(key, '')}"
    for col in "ABCDEFGHI":
        c = ws[f"{col}{HR}"]
        if col in headers:
            c.value = headers[col]
        c.fill = head_fill
        c.font = Font(name=B_FONT, size=10, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(left=Side(style="thin", color=OLIVE_DARK), right=Side(style="thin", color=OLIVE_DARK))

    # ---- Ditët ----
    day_fill = PatternFill("solid", fgColor=SAGE)
    zebra = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor=CREAM)]
    first = HR + 1
    for i, day in enumerate(DAYS):
        r = first + i
        ws.merge_cells(f"A{r}:B{r}")
        ws[f"A{r}"] = day.title()
        ws[f"A{r}"].font = Font(name=H_FONT, size=13, bold=True, color=OLIVE_DARK)
        ws[f"A{r}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in "AB":
            ws[f"{col}{r}"].fill = day_fill
            ws[f"{col}{r}"].border = grid
        lines = 3
        dm = plan["ushqimet"].get(day, {})
        for col, (key, _, _) in zip(meal_cols, MEALS):
            text = (dm.get(key) or "").strip()
            c = ws[f"{col}{r}"]
            c.value = text or None
            c.font = Font(name=B_FONT, size=10.5, color=TEXT)
            c.fill = zebra[i % 2]
            c.border = grid
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
            lines = max(lines, _est_lines(text, widths[col] - 2))
        ws.row_dimensions[r].height = min(409, max(58, lines * 14.5))
    last = first + len(DAYS) - 1

    ws.merge_cells(f"I{first}:I{last}")
    ws[f"I{first}"] = (plan.get("lista") or "").strip() or None
    ws[f"I{first}"].font = Font(name=B_FONT, size=10.5, color=TEXT)
    ws[f"I{first}"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
    for r in range(first, last + 1):
        ws[f"I{r}"].fill = PatternFill("solid", fgColor=CREAM)
        ws[f"I{r}"].border = grid

    # ---- Shënime ----
    nr = last + 2
    ws[f"A{nr}"] = "Shënime"
    ws[f"A{nr}"].font = Font(name=H_FONT, size=14, color=OLIVE)
    notes = (plan.get("shenime") or "").strip()
    n0, n1 = nr + 1, nr + 4
    ws.merge_cells(f"A{n0}:G{n1}")
    ws[f"A{n0}"] = notes or None
    ws[f"A{n0}"].font = Font(name=B_FONT, size=10.5, color=TEXT)
    ws[f"A{n0}"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
    for r in range(n0, n1 + 1):
        for col in "ABCDEFG":
            ws[f"{col}{r}"].fill = PatternFill("solid", fgColor=CREAM)
        ws[f"A{r}"].border = Border(left=Side(style="thick", color=OLIVE))
    note_lines = _est_lines(notes, sum(widths[c] for c in "ABCDEFG"))
    for r in range(n0, n1 + 1):
        ws.row_dimensions[r].height = max(16, note_lines * 14.5 / 4)

    # ---- Nënshkrimi ----
    ws.merge_cells(f"H{n0}:I{n0}")
    ws[f"H{n0}"] = "Përgatitur nga"
    ws[f"H{n0}"].font = Font(name=B_FONT, size=9, bold=True, color=MUTED)
    ws[f"H{n0}"].alignment = Alignment(horizontal="right")
    ws.merge_cells(f"H{n0 + 1}:I{n0 + 1}")
    ws[f"H{n0 + 1}"] = (plan.get("pergatitur_nga") or "").strip() or None
    ws[f"H{n0 + 1}"].font = Font(name=H_FONT, size=13, color=OLIVE_DARK)
    ws[f"H{n0 + 1}"].alignment = Alignment(horizontal="right", wrap_text=True)
    ws.merge_cells(f"H{n1}:I{n1}")
    ws[f"H{n1}"] = BRAND_LINE
    ws[f"H{n1}"].font = Font(name=B_FONT, size=8, color=OLIVE)
    ws[f"H{n1}"].alignment = Alignment(horizontal="right", vertical="bottom")

    # ---- Printimi ----
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_area = f"A1:I{n1}"
    ws.oddFooter.center.text = "Nora Limani Bektashi · Pri Nutrition"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = OLIVE

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =================================================================
# PDF
# =================================================================
_fonts_ok = None


def _register_fonts():
    """Regjistron Gilda Display + Lato; nëse mungojnë, bie te Times/Helvetica."""
    global _fonts_ok
    if _fonts_ok is not None:
        return _fonts_ok
    try:
        from reportlab.lib.fonts import addMapping
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        pdfmetrics.registerFont(TTFont("Gilda", os.path.join(FONTS, "GildaDisplay-Regular.ttf")))
        pdfmetrics.registerFont(TTFont("Lato", os.path.join(FONTS, "Lato-Regular.ttf")))
        pdfmetrics.registerFont(TTFont("Lato-Bold", os.path.join(FONTS, "Lato-Bold.ttf")))
        pdfmetrics.registerFont(TTFont("Lato-Italic", os.path.join(FONTS, "Lato-Italic.ttf")))
        addMapping("Lato", 0, 0, "Lato")
        addMapping("Lato", 1, 0, "Lato-Bold")
        addMapping("Lato", 0, 1, "Lato-Italic")
        addMapping("Lato", 1, 1, "Lato-Bold")
        _fonts_ok = True
    except Exception:
        _fonts_ok = False
    return _fonts_ok


def build_branded_pdf(plan) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (Image, KeepTogether, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    ok = _register_fonts()
    SERIF = "Gilda" if ok else "Times-Roman"
    SANS = "Lato" if ok else "Helvetica"
    SANS_B = "Lato-Bold" if ok else "Helvetica-Bold"
    C = lambda h: colors.HexColor("#" + h)  # noqa: E731

    def st(name, **kw):
        base = dict(fontName=SANS, fontSize=8, leading=10, textColor=C(TEXT))
        base.update(kw)
        return ParagraphStyle(name, **base)

    title_st = st("title", fontName=SERIF, fontSize=22, leading=24, textColor=C(OLIVE), alignment=TA_RIGHT)
    sub_st = st("sub", fontName=SANS_B, fontSize=9, leading=12, textColor=C(MUTED), alignment=TA_RIGHT)
    lab_st = st("lab", fontName=SANS_B, fontSize=6.5, leading=8, textColor=C(MUTED))
    val_st = st("val", fontName=SERIF, fontSize=12, leading=14)
    head_st = st("head", fontName=SANS_B, fontSize=7.5, leading=9, textColor=colors.white, alignment=TA_CENTER)
    day_st = st("day", fontName=SERIF, fontSize=10, leading=12, textColor=C(OLIVE_DARK), alignment=TA_CENTER)
    cell_st = st("cell", fontSize=6.7, leading=7.8)
    sec_st = st("sec", fontName=SERIF, fontSize=12, leading=14, textColor=C(OLIVE))
    body_st = st("body", fontSize=8, leading=10.5)
    sign_lab = st("signl", fontName=SANS_B, fontSize=6.5, leading=8, textColor=C(MUTED), alignment=TA_RIGHT)
    sign_st = st("sign", fontName=SERIF, fontSize=12, leading=14, textColor=C(OLIVE_DARK), alignment=TA_RIGHT)

    def P(text, style):
        return Paragraph(escape(text or "").replace("\n", "<br/>"), style)

    page_w, page_h = landscape(A4)
    margin = 1.0 * cm
    avail = page_w - 2 * margin

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(C(LINE))
        canvas.setLineWidth(0.5)
        canvas.line(margin, 0.75 * cm, page_w - margin, 0.75 * cm)
        canvas.setFont(SANS, 6.5)
        canvas.setFillColor(C(OLIVE))
        canvas.drawString(margin, 0.45 * cm, BRAND_LINE)
        canvas.drawRightString(page_w - margin, 0.45 * cm, f"{plan['emri']}  ·  {plan['titulli']}  ·  faqe {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=(page_w, page_h), leftMargin=margin, rightMargin=margin,
                            topMargin=0.5 * cm, bottomMargin=1.0 * cm,
                            title=f"Plan ushqimor - {plan['emri']}", author=plan.get("pergatitur_nga") or "")
    story = []

    # ---- Koka: logo + titulli ----
    if os.path.exists(LOGO):
        from PIL import Image as PILImage
        w, h = PILImage.open(LOGO).size
        logo_h = 1.6 * cm
        logo = Image(LOGO, width=logo_h * w / h, height=logo_h)
    else:
        logo = P("Nora Limani Bektashi", st("lf", fontName=SERIF, fontSize=16, textColor=C(OLIVE)))
    right = [P("Plan Ushqimor", title_st), P(f"{plan['titulli']}   ·   {_period(plan)}", sub_st)]
    head = Table([[logo, right]], colWidths=[avail * 0.5, avail * 0.5])
    head.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "BOTTOM"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
        ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("LINEBELOW", (0, 0), (-1, -1), 1.2, C(OLIVE)),
    ]))
    story += [head, Spacer(1, 0.15 * cm)]

    # ---- Klienti ----
    info = Table([[P("KLIENTI", lab_st), P("PESHA", lab_st), P("GJATËSIA", lab_st), P("BMI", lab_st)],
                  [P(plan["emri"], val_st), P(fmt_num(plan["pesha"], "kg"), val_st),
                   P(fmt_num(plan["gjatesia"], "cm"), val_st), P(_bmi(plan), val_st)]],
                 colWidths=[avail * 0.40, avail * 0.15, avail * 0.15, avail * 0.30])
    info.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
    ]))
    story += [info, Spacer(1, 0.2 * cm)]

    # ---- Tabela javore ----
    rows = [[P("", head_st)] + [P(f"{n.upper()}\n{plan['oraret'].get(k, '')}", head_st) for k, n, _ in MEALS]]
    for day in DAYS:
        dm = plan["ushqimet"].get(day, {})
        rows.append([P(day.title(), day_st)] + [P((dm.get(k) or "").strip(), cell_st) for k, _, _ in MEALS])
    first = 2.3 * cm
    weights = [1.6, 0.9, 1.8, 0.9, 1.6, 0.9]
    cw = [first] + [(avail - first) * w / sum(weights) for w in weights]
    t = Table(rows, colWidths=cw, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), C(OLIVE)),
        ("BACKGROUND", (0, 1), (0, -1), C(SAGE)),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, C(LINE)),
        ("LINEAFTER", (0, 0), (-2, -1), 0.4, C(LINE)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ("LEFTPADDING", (1, 1), (-1, -1), 5), ("RIGHTPADDING", (1, 1), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, 0), 4), ("BOTTOMPADDING", (0, 0), (-1, 0), 4),
    ]
    for r in range(1, len(rows)):
        if r % 2 == 0:
            style.append(("BACKGROUND", (1, r), (-1, r), C(CREAM)))
    t.setStyle(TableStyle(style))
    story.append(t)

    # ---- Shënime / Lista / Nënshkrimi ----
    left = []
    if (plan.get("shenime") or "").strip():
        left += [P("Shënime", sec_st), Spacer(1, 2), P(plan["shenime"].strip(), body_st)]
    if (plan.get("lista") or "").strip():
        left += [Spacer(1, 6), P("Lista ushqimore", sec_st), Spacer(1, 2), P(plan["lista"].strip(), body_st)]
    sign = [P("PËRGATITUR NGA", sign_lab), Spacer(1, 2),
            P((plan.get("pergatitur_nga") or "").strip(), sign_st)]
    if left or (plan.get("pergatitur_nga") or "").strip():
        box = Table([[left or [P("", body_st)], sign]], colWidths=[avail * 0.68, avail * 0.32])
        box_style = [("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("LEFTPADDING", (0, 0), (0, 0), 9), ("RIGHTPADDING", (1, 0), (1, 0), 0),
                     ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]
        if left:
            box_style += [("BACKGROUND", (0, 0), (0, 0), C(CREAM)),
                          ("LINEBEFORE", (0, 0), (0, 0), 2.5, C(OLIVE))]
        box.setStyle(TableStyle(box_style))
        story += [Spacer(1, 0.25 * cm), KeepTogether(box)]

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()
