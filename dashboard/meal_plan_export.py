# -*- coding: utf-8 -*-
"""
meal_plan_export.py
=====================
Gjeneron planin ushqimor javor (7 ditë x 6 vakte) në Excel, sipas
modelit të dietologes (Plan_ushqimor_3_XB.xlsx), dhe në PDF për printim.

`plan` është dict me strukturën që ndërton page_meal_plan._collect_plan():
    emri, pesha, gjatesia, titulli, data_fillimit, data_mbarimit (date),
    oraret {meal_key: "8:00"}, ushqimet {day: {meal_key: tekst}},
    shenime, lista, pergatitur_nga
"""
import io
import math
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

DAYS = ["E HËNË", "E MARTË", "E MËRKURË", "E ENJTE", "E PREMTE", "E SHTUNË", "E DIEL"]
MEALS = [  # (çelësi, emri, ora e paracaktuar -- sipas modelit origjinal)
    ("mengjesi", "Mëngjesi", "8:00"),
    ("mesvakt1", "Mesvakt", "10:00"),
    ("dreka", "Dreka", "12:00"),
    ("mesvakt2", "Mesvakt", "14:00"),
    ("darka", "Darka", "18:00"),
    ("mesvakt3", "Mesvakt", "21:00"),
]

# Ngjyrat e modelit origjinal (tema "Accent 6" jeshile e Office)
GREEN_DARK = "375623"
GREEN_MID = "A9D08E"
GREEN_LIGHT = "E2EFDA"
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def fmt_date(d):
    """8.05.2023 -- formati i përdorur në modelin origjinal."""
    return f"{d.day}.{d.month:02d}.{d.year}" if d else "..."


def fmt_num(x, unit):
    if x is None or x == 0:
        return f"... {unit}"
    return f"{x:g} {unit}"


def initials(name):
    parts = [p for p in (name or "").split() if p]
    return "".join(p[0].upper() for p in parts) or "KLIENT"


def title_text(plan):
    return (
        "PLAN USHQIMOR\n"
        f"{plan['titulli']} ({fmt_date(plan['data_fillimit'])} - {fmt_date(plan['data_mbarimit'])})\n"
        f"{plan['emri']}\n"
        f"Pesha: {fmt_num(plan['pesha'], 'kg')} ; Gjatësia: {fmt_num(plan['gjatesia'], 'cm')}"
    )


def _est_lines(text, col_width):
    """Vlerëson sa rreshta zë teksti në një qelizë (për lartësinë e rreshtit)."""
    if not text:
        return 1
    chars_per_line = max(8, int(col_width * 1.15))
    return sum(max(1, math.ceil(len(line) / chars_per_line)) for line in str(text).split("\n"))


def build_meal_plan_xlsx(plan) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Plani ushqimor"

    widths = {"A": 15, "B": 17.3, "C": 43.3, "D": 22, "E": 50.4, "F": 20.6, "G": 41.7, "H": 26.6, "I": 34.9}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    meal_cols = ["C", "D", "E", "F", "G", "H"]
    base_font = Font(name="Arial", size=11)

    # ---- Titulli (A1:H1) ----
    ws.merge_cells("A1:H1")
    ws["A1"] = title_text(plan)
    ws["A1"].font = Font(name="Arial", size=22)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 125

    # ---- Koka e vakteve (rreshti 2) ----
    ws.merge_cells("A2:B2")
    header_fill = PatternFill("solid", fgColor=GREEN_DARK)
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    for col in ["A", "B"]:
        ws[f"{col}2"].fill = header_fill
        ws[f"{col}2"].border = BORDER
    for col, (key, name, _) in zip(meal_cols, MEALS):
        c = ws[f"{col}2"]
        c.value = f"{name}\n({plan['oraret'].get(key, '')})"
        c.font, c.fill, c.border = header_font, header_fill, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c = ws["I2"]
    c.value = "Lista ushqimore"
    c.font, c.fill, c.border = header_font, header_fill, BORDER
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 42.75

    # ---- Ditët (rreshtat 3-9) ----
    day_fill = PatternFill("solid", fgColor=GREEN_MID)
    cell_fill = PatternFill("solid", fgColor=GREEN_LIGHT)
    for i, day in enumerate(DAYS):
        r = 3 + i
        ws.merge_cells(f"A{r}:B{r}")
        ws[f"A{r}"] = day
        ws[f"A{r}"].font = Font(name="Arial", size=12, bold=True)
        ws[f"A{r}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col in ["A", "B"]:
            ws[f"{col}{r}"].fill = day_fill
            ws[f"{col}{r}"].border = BORDER

        lines = 3
        day_meals = plan["ushqimet"].get(day, {})
        for col, (key, _, _) in zip(meal_cols, MEALS):
            text = (day_meals.get(key) or "").strip()
            c = ws[f"{col}{r}"]
            c.value = text or None
            c.font, c.fill, c.border = base_font, cell_fill, BORDER
            c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            lines = max(lines, _est_lines(text, widths[col]))
        ws.row_dimensions[r].height = min(409, max(60, lines * 14.5))

    # ---- Lista ushqimore (I3:I9) ----
    ws.merge_cells("I3:I9")
    ws["I3"] = (plan.get("lista") or "").strip() or None
    ws["I3"].font = base_font
    ws["I3"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for r in range(3, 10):
        ws[f"I{r}"].fill = cell_fill
        ws[f"I{r}"].border = BORDER

    # ---- Shënime (B11, B12:H15) ----
    ws["B11"] = "SHENIME:"
    ws["B11"].font = Font(name="Arial", size=11, bold=True)
    ws.merge_cells("B12:H15")
    notes = (plan.get("shenime") or "").strip()
    ws["B12"] = notes or None
    ws["B12"].font = base_font
    ws["B12"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    note_lines = _est_lines(notes, sum(widths[c] for c in "BCDEFGH"))
    per_row = max(15, note_lines * 14.5 / 4)
    for r in range(12, 16):
        ws.row_dimensions[r].height = per_row

    # ---- Përgatitur nga (I18, I19) ----
    ws["I18"] = "Përgatitur nga:"
    ws["I18"].font = base_font
    ws["I18"].alignment = Alignment(horizontal="right", vertical="center")
    ws["I19"] = (plan.get("pergatitur_nga") or "").strip() or None
    ws["I19"].font = Font(name="Arial", size=11, bold=True)
    ws["I19"].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # ---- Printimi: landscape, 1 faqe në gjerësi ----
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:I19"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_meal_plan_pdf(plan) -> bytes:
    """Version PDF (landscape) i të njëjtit plan -- për printim/dërgim te klienti."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()
    title_st = ParagraphStyle("t", parent=ss["Normal"], fontSize=12.5, leading=15, alignment=TA_CENTER)
    head_st = ParagraphStyle("h", parent=ss["Normal"], fontSize=8, leading=10, alignment=TA_CENTER,
                             textColor=colors.white, fontName="Helvetica-Bold")
    day_st = ParagraphStyle("d", parent=ss["Normal"], fontSize=8, leading=10, alignment=TA_CENTER,
                            fontName="Helvetica-Bold")
    cell_st = ParagraphStyle("c", parent=ss["Normal"], fontSize=6.8, leading=8)
    norm = ParagraphStyle("n", parent=ss["Normal"], fontSize=8.5, leading=11)

    def P(text, st):
        return Paragraph(escape(text or "").replace("\n", "<br/>"), st)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=1 * cm, rightMargin=1 * cm,
                            topMargin=0.6 * cm, bottomMargin=0.6 * cm)
    story = [P(title_text(plan), title_st), Spacer(1, 0.3 * cm)]

    header = [P("", head_st)] + [P(f"{n}\n({plan['oraret'].get(k, '')})", head_st) for k, n, _ in MEALS]
    rows = [header]
    for day in DAYS:
        dm = plan["ushqimet"].get(day, {})
        rows.append([P(day, day_st)] + [P((dm.get(k) or "").strip(), cell_st) for k, _, _ in MEALS])

    total_w = landscape(A4)[0] - 2 * cm
    first = 2.2 * cm
    weights = [1.6, 0.9, 1.8, 0.9, 1.6, 0.9]
    col_w = [first] + [(total_w - first) * w / sum(weights) for w in weights]
    t = Table(rows, colWidths=col_w, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + GREEN_DARK)),
        ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#" + GREEN_MID)),
        ("BACKGROUND", (1, 1), (-1, -1), colors.HexColor("#" + GREEN_LIGHT)),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(t)

    if (plan.get("lista") or "").strip():
        story += [Spacer(1, 0.3 * cm), Paragraph("<b>Lista ushqimore:</b>", norm), P(plan["lista"], norm)]
    if (plan.get("shenime") or "").strip():
        story += [Spacer(1, 0.3 * cm), Paragraph("<b>SHENIME:</b>", norm), P(plan["shenime"], norm)]
    if (plan.get("pergatitur_nga") or "").strip():
        right = ParagraphStyle("r", parent=norm, alignment=2)
        story += [Spacer(1, 0.4 * cm), Paragraph("Përgatitur nga:", right),
                  Paragraph(f"<b>{escape(plan['pergatitur_nga'])}</b>", right)]

    doc.build(story)
    return buf.getvalue()
